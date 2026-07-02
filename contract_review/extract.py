"""從 .docx 與 .pdf 檔案抽取文字段落，並保留可回報的位置資訊。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class Segment:
    """一段可供規則檢查的文字，location 用於報告中標示出處。"""

    text: str
    location: str


def extract_segments(path: str | Path) -> list[Segment]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".doc":
        raise ValueError(
            f"不支援舊版 .doc 格式：{path.name}。請先在 Word 中另存為 .docx 再審查。"
        )
    raise ValueError(f"不支援的檔案格式：{path.suffix}（僅支援 .docx 與 .pdf）")


def _extract_docx(path: Path) -> list[Segment]:
    import docx

    document = docx.Document(str(path))
    segments: list[Segment] = []
    for i, para in enumerate(document.paragraphs, start=1):
        text = para.text.strip()
        if text:
            segments.append(Segment(text=text, location=f"段落 {i}"))
    for t_idx, table in enumerate(document.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            text = "　".join(c for c in cells if c)
            if text:
                segments.append(
                    Segment(text=text, location=f"表格 {t_idx} 第 {r_idx} 列")
                )
    return segments


def _extract_pdf(path: Path) -> list[Segment]:
    import pdfplumber

    segments: list[Segment] = []
    with pdfplumber.open(str(path)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            for line_no, line in enumerate(text.splitlines(), start=1):
                line = line.strip()
                if line:
                    segments.append(
                        Segment(text=line, location=f"第 {page_no} 頁第 {line_no} 行")
                    )
    if not segments:
        raise ValueError(
            f"{path.name} 無法抽取任何文字，可能是掃描影像 PDF。"
            "請先經過 OCR 轉成文字型 PDF 再審查。"
        )
    return segments
